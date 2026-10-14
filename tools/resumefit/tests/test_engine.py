from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import re

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_COLOR_INDEX
from lxml import etree

from app.engine import edit_docx, find_candidates, validate_docx


def example_document(path: Path):
    doc = Document()
    sec=doc.sections[0]
    sec.top_margin=Inches(.6)
    sec.bottom_margin=Inches(.6)
    head=doc.add_paragraph()
    name=head.add_run('ALEX MORGAN')
    name.bold=True
    name.font.size=Pt(18)
    doc.add_paragraph('alex@example.com | Boston, MA')
    doc.add_paragraph('SUMMARY')
    p=doc.add_paragraph('Backend engineer with 6 years of experience building APIs and maintaining database-backed systems with Python and SQL.')
    p.style='Normal'
    doc.add_paragraph('SKILLS')
    tab=doc.add_table(rows=1,cols=1)
    tab.cell(0,0).paragraphs[0].text='Python, SQL, REST APIs, Docker, PostgreSQL, AWS'
    doc.add_paragraph('EXPERIENCE')
    doc.add_paragraph('Senior Engineer  |  2020–Present  |  Example Company')
    p=doc.add_paragraph(style='List Bullet')
    p.add_run('Built ')
    bold=p.add_run('REST APIs')
    bold.bold=True
    p.add_run(' with Python and PostgreSQL supporting 25 business teams and daily data workflows.')
    p=doc.add_paragraph(style='List Bullet')
    p.add_run('Improved deployment pipelines with Docker and GitHub Actions, reducing manual release effort.')
    doc.add_paragraph('EDUCATION')
    doc.add_paragraph('B.S. Computer Science | 2017')
    doc.save(path)


def _scrub_text(doc: Document):
    element=deepcopy(doc.element.body)
    for t in element.xpath('.//w:t'):
        t.text='X'
        t.attrib.pop('{http://www.w3.org/XML/1998/namespace}space', None)
    return etree.tostring(element, method='c14n')


def test_docx_preserves_every_structure_except_text(tmp_path):
    source=tmp_path/'original.docx'
    target=tmp_path/'tailored.docx'
    example_document(source)
    validate_docx(source.read_bytes())
    original=Document(source)
    candidates=find_candidates(original)
    assert {'summary', 'skills', 'bullet'} <= {c.category for c in candidates}
    assert len(candidates)==4
    skill=next(c for c in candidates if c.category=='skills')
    bullet=next(c for c in candidates if c.category=='bullet')
    changes={
       skill.id: 'AWS, PostgreSQL, Docker, Python, SQL, REST APIs',
       bullet.id: 'Built scalable REST APIs using Python and PostgreSQL supporting 25 business teams and daily workflows.',
    }
    edit_docx(source,target,changes,candidates)
    output=Document(target)
    assert _scrub_text(original)==_scrub_text(output)
    assert 'ALEX MORGAN' in '\n'.join(x.text for x in output.paragraphs)
    assert 'B.S. Computer Science | 2017' in '\n'.join(x.text for x in output.paragraphs)
    assert output.tables[0].cell(0,0).text == changes[skill.id]
    assert 'supporting 25' in '\n'.join(x.text for x in output.paragraphs)
    assert '2020–Present' in '\n'.join(x.text for x in output.paragraphs)
    assert any(run.bold and run.text == 'REST APIs' for para in output.paragraphs for run in para.runs)


def test_invalid_changes_rejected(tmp_path):
    source=tmp_path/'in.docx';example_document(source)
    try:
        edit_docx(source,tmp_path/'out.docx',{'p9999':'Fake job'})
    except ValueError as exc:
        assert 'not an allowed editable' in str(exc)
    else:
        raise AssertionError('Unexpected successful edit')


def test_manual_inline_label_is_locked(tmp_path):
    source=tmp_path/'in.docx'
    doc=Document()
    doc.add_paragraph('Skills: Python, Java, React, SQL')
    doc.save(source)
    options=find_candidates(Document(source))
    assert len(options)==1
    assert options[0].locked_prefix=='Skills: '
    edit_docx(source,tmp_path/'out.docx',{options[0].id:'SQL, Java, React, Python'},options)
    assert Document(tmp_path/'out.docx').paragraphs[0].text=='Skills: SQL, Java, React, Python'


def test_education_bullets_remain_protected(tmp_path):
    src=tmp_path/'protect.docx'
    doc=Document()
    doc.add_paragraph('EDUCATION')
    doc.add_paragraph('Earned 4.0 GPA and graduated in Computer Science in 2020.',style='List Bullet')
    doc.add_paragraph('WORK EXPERIENCE')
    doc.add_paragraph('Built user interfaces for enterprise applications with JavaScript and CSS.',style='List Bullet')
    doc.save(src)
    options=find_candidates(Document(src))
    assert len(options)==1
    assert options[0].category=='bullet'
    assert 'Built user interfaces' in options[0].original
