"""Tests avoid network calls and assert that letters use only real resume evidence."""
from types import SimpleNamespace
import sys

import pytest
from docx import Document

from app import cover_letter

RESUME = '''ALEX MORGAN
alex@example.com
Professional Summary
Software engineer with 5 years developing Java Spring Boot backend APIs.
Experience
Created banking payment services using Spring Boot and PostgreSQL.
Improved API response time by 35%.
Built AI-assisted developer tooling in Python.
'''
JD = '''What You'll Do
- Build reliable backend systems for financial workflows
- Build analytics APIs and process operational data
- Implement AI-powered document workflows
What We're Looking For
- Experience with data-intensive applications and backend services
'''


def fake_response(parts):
    import json
    return SimpleNamespace(choices=[SimpleNamespace(finish_reason='stop', message=SimpleNamespace(content=json.dumps({'paragraphs': parts})))])


def test_groq_uses_jd_and_original_resume_evidence(monkeypatch):
    observed = {}
    parts = [
        'I am applying for the Senior Platform Engineer role at Confido, bringing experience building backend financial services.',
        'I developed banking payment APIs with Java, Spring Boot and PostgreSQL, improved API response time by 35%, and built AI-assisted tooling with Python.',
        'I would welcome a conversation about applying this backend and financial-platform experience to reliable software for your team.',
    ]
    class FakeOpenAI:
        def __init__(self, **kwargs):
            observed['client'] = kwargs
            def create(**opts):
                observed['options'] = opts
                return fake_response(parts)
            self.chat = SimpleNamespace(completions=SimpleNamespace(create=create))
    monkeypatch.setenv('GROQ_API_KEY','gsk-test')
    monkeypatch.delenv('OPENAI_API_KEY',raising=False)
    monkeypatch.setitem(sys.modules,'openai',SimpleNamespace(OpenAI=FakeOpenAI))
    answer = cover_letter.generate(RESUME, JD, 'Confido', 'Senior Platform Engineer')
    assert 'Dear Confido Hiring Team,' in answer
    assert '35%' in answer
    assert answer.endswith('Sincerely,\nALEX MORGAN')
    assert observed['client']['base_url'] == 'https://api.groq.com/openai/v1'
    assert observed['options']['model'] == 'openai/gpt-oss-20b'
    assert 'actual_resume_evidence' in observed['options']['messages'][1]['content']
    assert 'job_requirements' in observed['options']['messages'][1]['content']


def test_no_key_returns_clear_error(monkeypatch):
    monkeypatch.delenv('GROQ_API_KEY',raising=False)
    monkeypatch.delenv('OPENAI_API_KEY',raising=False)
    with pytest.raises(ValueError,match='Groq API key'):
        cover_letter.generate(RESUME, JD)


def test_does_not_fabricate_experience_or_metrics():
    paras=['I built fintech systems that help people access core services on a reliable platform.',
           'I have delivered backend APIs in Spring Boot, with experience improving performance by 92%.',
           'I would bring relevant experience to the engineering team and welcome a conversation about the role.']
    with pytest.raises(ValueError,match='numerical claims'):
        cover_letter._sanitize_paragraphs({'paragraphs':paras},RESUME)


def test_requires_three_reasonable_paragraphs():
    with pytest.raises(ValueError,match='incomplete'):
        cover_letter._sanitize_paragraphs({'paragraphs':['One sentence']},RESUME)


def test_docx_contains_correct_name_and_company(tmp_path):
    path=tmp_path/'letter.docx'
    cover_letter.make_docx(path, 'Dear Hiring Team,\n\nI have experience in backend systems.\n\nSincerely,\nALEX MORGAN',RESUME,'Confido','Platform Engineer')
    doc=Document(path)
    all_text='\n'.join(p.text for p in doc.paragraphs)
    assert 'ALEX MORGAN' in all_text
    assert 'alex@example.com' in all_text
    assert 'Platform Engineer' in all_text
    assert 'Confido' in all_text
    assert 'Dear Hiring Team' in all_text


def test_cover_letter_api_uses_approved_resume_edits_and_exports(monkeypatch, tmp_path):
    from fastapi.testclient import TestClient
    from app import main
    from app.engine import find_candidates
    from io import BytesIO
    monkeypatch.setattr(main, 'SESSIONS', tmp_path)
    monkeypatch.delenv('RESUMEFIT_REQUIRE_AUTH', raising=False)
    monkeypatch.delenv('RESUMEFIT_ACCESS_PASSWORD', raising=False)
    doc = Document()
    doc.add_paragraph('ALEX MORGAN')
    doc.add_paragraph('alex@example.com')
    doc.add_paragraph('SUMMARY')
    doc.add_paragraph('Engineer focused on Java REST APIs for finance platforms.')
    doc.add_paragraph('EXPERIENCE')
    doc.add_paragraph('Built Java services for secure payment processing.', style='List Bullet')
    data = BytesIO(); doc.save(data)
    captured = {}
    def fake_generate(resume_text, jd, company, role):
        captured['resume_text'] = resume_text
        captured['jd'] = jd
        return 'Dear Hiring Team,\n\nI build Java APIs for secure financial workflows.\n\nSincerely,\nALEX MORGAN'
    monkeypatch.setattr(main, 'generate_cover_letter', fake_generate)
    with TestClient(main.app) as client:
        up = client.post('/api/analyze', files={'file':('resume.docx',data.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
        assert up.status_code == 200, up.text
        info = up.json()
        assert info['candidates']
        bullet = next(c for c in info['candidates'] if c['category'] == 'bullet')
        update = {'id':bullet['id'],'text':'Developed reliable Java services for secure payment processing.'}
        payload = {'session_id':info['session_id'],'jd':JD,'company':'Confido','role':'Senior Platform Engineer','lines':[update]}
        draft = client.post('/api/cover-letter',json=payload)
        assert draft.status_code == 200, draft.text
        assert 'Developed reliable Java services' in captured['resume_text']
        assert draft.json()['letter_text'].startswith('Dear Hiring Team')
        export = client.post('/api/cover-letter/export',json={'session_id':info['session_id'],'letter_text':draft.json()['letter_text'],'company':'Confido','role':'Senior Platform Engineer','output_format':'docx'})
        assert export.status_code == 200, export.text
        assert export.content.startswith(b'PK')
        got=Document(BytesIO(export.content))
        assert any('Confido' in p.text for p in got.paragraphs)
        assert any('ALEX MORGAN' in p.text for p in got.paragraphs)


def test_cover_letter_endpoint_rejects_invalid_session(monkeypatch, tmp_path):
    from fastapi.testclient import TestClient
    from app import main
    monkeypatch.setattr(main, 'SESSIONS', tmp_path)
    with TestClient(main.app) as client:
        response=client.post('/api/cover-letter',json={'session_id':'0'*32,'jd':JD})
        assert response.status_code == 404
