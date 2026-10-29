"""Transparent before/after explanations of genuinely proposed resume changes."""
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from app.change_log import describe_changes
from app.engine import Candidate
from app import main


def candidate(ident: str, category: str, value: str) -> Candidate:
    return Candidate(id=ident, index=1, category=category,
                     original=value, editable_original=value,
                     locked_prefix='', section=category, style='Normal')


JD = """Responsibilities
- Build secure backend REST APIs for banking payment processing
- Automate deployments and maintain CI/CD engineering workflows
- Use Java and Spring Boot for enterprise microservices
"""


def test_details_include_real_before_after_and_grounded_jd_requirement():
    source = 'Built Java REST APIs for card processing.'
    edited = 'Built secure Java REST APIs for banking payment processing.'
    rows = describe_changes([candidate('p3', 'bullet', source)], {'p3': edited}, JD)
    assert rows['p3']['original'] == source
    assert rows['p3']['revised'] == edited
    assert rows['p3']['category'] == 'bullet'
    assert 'secure' in rows['p3']['what_changed']
    assert 'banking payment processing' in rows['p3']['jd_requirement'].lower()
    assert 'ATS' not in rows['p3']['why']
    assert 'verify' in rows['p3']['why'].lower()


def test_unchanged_edits_and_unknown_candidate_ids_are_excluded():
    original = 'Built Java REST APIs for payment platforms.'
    assert describe_changes([candidate('p5', 'summary', original)], {
        'p5': original, 'p999': 'Invented experiences for a fictional job'
    }, JD) == {}


def test_skill_prioritization_explains_reorder_without_claiming_new_skills():
    original = 'Languages: Python | Java | SQL | TypeScript'
    selected = 'Languages: Java | SQL | Python | TypeScript'
    rows = describe_changes([candidate('p7', 'skills', original)], {'p7': selected}, JD)
    assert rows['p7']['original'] == original
    assert rows['p7']['revised'] == selected
    assert 'no skills were added or removed' in rows['p7']['what_changed']


def test_tailor_endpoint_returns_the_change_details_from_actual_proposed_edits(tmp_path, monkeypatch):
    monkeypatch.setattr(main, 'SESSIONS', tmp_path)
    monkeypatch.delenv('RESUMEFIT_ACCESS_PASSWORD', raising=False)
    monkeypatch.delenv('RESUMEFIT_REQUIRE_AUTH', raising=False)
    doc = Document()
    doc.add_paragraph('SUMMARY')
    doc.add_paragraph('Backend engineer building REST APIs for payments.')
    doc.add_paragraph('PROFESSIONAL EXPERIENCE')
    doc.add_paragraph('Built Java REST APIs for card processing.', style='List Bullet')
    raw = BytesIO()
    doc.save(raw)
    with TestClient(main.app) as client:
        uploaded = client.post('/api/analyze', files={'file': (
            'original.docx', raw.getvalue(),
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
        assert uploaded.status_code == 200, uploaded.text
        item = uploaded.json()
        bullet = next(c for c in item['candidates'] if c['category'] == 'bullet')
        original = bullet['editable_original']
        improved = 'Built secure Java REST APIs for banking card processing.'
        def mock_suggest(candidates, jd, full_original_resume):
            assert original in full_original_resume
            return {bullet['id']: improved}, 'groq', []
        monkeypatch.setattr(main, 'suggest', mock_suggest)
        response = client.post('/api/tailor', json={
            'session_id': item['session_id'], 'jd': JD
        })
        assert response.status_code == 200, response.text
        data = response.json()
        assert data['suggestions'][bullet['id']] == improved
        detail = data['change_details'][bullet['id']]
        assert detail['original'] == original
        assert detail['revised'] == improved
        assert detail['jd_requirement']
        assert (tmp_path / item['session_id'] / 'source.docx').read_bytes() == raw.getvalue()
