"""ATS-style estimator: keyword evidence only, no AI or fabricated employer ATS scores."""
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from app.ats import estimate
from app import main

JD = """What You'll Do
- Build Java Spring Boot APIs for financial payment workflows
- Build event ingestion pipelines using Kafka and Spark
- Deliver AI-powered document processing using Python
"""


def test_different_resumes_get_different_estimated_coverage():
    weak = 'Graphic designer working with digital illustrations and layout design.'
    relevant = ('Built Java Spring Boot APIs for financial payment workflows.\n'
                'Built Kafka and Spark ingestion pipelines.\n'
                'Delivered AI-powered document processing using Python.')
    poor = estimate(JD, weak)
    better = estimate(JD, relevant)
    assert poor['original_score'] < better['original_score']
    assert better['original_score'] >= 65
    assert 0 <= better['tailored_score'] <= 100
    assert better['evaluated'] == 3
    assert all(item['tailored_status'] == 'matched' for item in better['requirements'])
    assert better['original_score'] == better['tailored_score']


def test_proposed_edits_update_score_but_dont_claim_ats_prediction():
    original = 'Implemented Java APIs to support reliable customer services.'
    tailored = 'Implemented Java Spring Boot APIs for financial payment workflows.'
    result = estimate(JD, original, tailored)
    assert result['tailored_score'] > result['original_score']
    assert any(x['tailored_status'] == 'matched' for x in result['requirements'])
    assert all('evidence' in item for item in result['requirements'])


def test_keyword_repetition_does_not_game_matching():
    initial = 'Worked with Java and Spring Boot APIs.'
    spam = initial + ' Java Java Java Java Spring Spring Spring'
    jd = 'Responsibilities\n- Build Java Spring Boot APIs for banking systems'
    result = estimate(jd, initial, spam)
    assert result['tailored_score'] == result['original_score']


def test_scoring_never_introduces_resume_edits(tmp_path, monkeypatch):
    monkeypatch.setattr(main, 'SESSIONS', tmp_path)
    monkeypatch.delenv('RESUMEFIT_ACCESS_PASSWORD', raising=False)
    monkeypatch.delenv('RESUMEFIT_REQUIRE_AUTH', raising=False)
    d = Document()
    d.add_paragraph('ALEX MORGAN')
    d.add_paragraph('SUMMARY')
    d.add_paragraph('Software engineer working with Java and backend APIs.')
    d.add_paragraph('EXPERIENCE')
    d.add_paragraph('Built Java APIs for financial platforms.', style='List Bullet')
    d.add_paragraph('EDUCATION')
    d.add_paragraph('Bachelor of Science in Computer Science 2020')
    buffer = BytesIO()
    d.save(buffer)
    with TestClient(main.app) as client:
        res = client.post('/api/analyze', files={
            'file': ('source.docx', buffer.getvalue(),
                     'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        })
        assert res.status_code == 200, res.text
        item = res.json()
        original = (tmp_path / item['session_id'] / 'source.docx').read_bytes()
        body = {'session_id': item['session_id'], 'jd': JD, 'lines': []}
        baseline = client.post('/api/ats-score', json=body)
        assert baseline.status_code == 200, baseline.text
        assert baseline.json()['original_score'] == baseline.json()['tailored_score']
        summary = next(c for c in item['candidates'] if c['category'] == 'summary')
        body['lines'] = [{
            'id': summary['id'],
            'text': 'Java Spring Boot API engineer supporting financial payment workflows.'
        }]
        updated = client.post('/api/ats-score', json=body)
        assert updated.status_code == 200, updated.text
        assert updated.json()['original_score'] == baseline.json()['original_score']
        assert updated.json()['tailored_score'] >= baseline.json()['original_score']
        assert (tmp_path / item['session_id'] / 'source.docx').read_bytes() == original
        forged = client.post('/api/ats-score', json={
            'session_id': item['session_id'], 'jd': JD,
            'lines': [{'id': 'p6', 'text': 'Invented education credentials'}]
        })
        assert forged.status_code == 422
        assert 'only approved resume paragraphs' in forged.text
        repeated = client.post('/api/ats-score', json={
            'session_id': item['session_id'], 'jd': JD, 'lines': [body['lines'][0]]*2
        })
        assert repeated.status_code == 400


def test_score_requires_a_real_session(tmp_path, monkeypatch):
    monkeypatch.setattr(main, 'SESSIONS', tmp_path)
    with TestClient(main.app) as client:
        missing = client.post('/api/ats-score', json={
            'session_id': 'a'*32, 'jd': JD, 'lines': []
        })
        assert missing.status_code == 404
